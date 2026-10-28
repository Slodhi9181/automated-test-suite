# -*- coding: utf-8 -*-
"""
sequence_panel.py — Manual test-sequence editor and runner.
"""

import json
import math
import time
import threading
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font as XFont, PatternFill, Alignment as XAlign

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QTableWidget, QTableWidgetItem, QHeaderView,
    QPushButton, QLabel, QTextEdit, QLineEdit, QFileDialog,
    QMessageBox, QProgressBar, QAbstractItemView,
    QMenu, QComboBox, QGroupBox, QDoubleSpinBox, QSpinBox,
    QCheckBox, QStyleOptionViewItem, QFrame, QSplitter,
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QTimer
from PyQt6.QtGui import QFont, QTextCursor, QColor, QBrush
from PyQt6.QtWidgets import QStyledItemDelegate

from workers import (use_instrument, shared_resource, shared_driver,
                     check_instruments, _INSTR_NAMES)
from test_automation import GWGPP4323, ADDR, WT5000Driver, ITechFixed, MDO34Driver
from sweep_panel import _mdo_waveform_png
from daq_espec_panel import daq_scan, _espec_setup, espec_set_temp, espec_set_humi
from ui_helpers import unique_path
from app_logger import get_logger, APP_DIR

_log = get_logger("sequence")

_STEPS_SESSION = APP_DIR / "sequence_session.json"


# ─────────────────────────────────────────────────────────────────────────────
# Column definitions  (timing moved to common group above table)
# ─────────────────────────────────────────────────────────────────────────────

_COLS = [
    "IT En", "ITECH V (V)", "ITECH I (A)",          # 0-2
    "G1 En", "G1 Mode",    "G1 V (V)",   "G1 I (A)",  # 3-6
    "G2 En", "G2 Mode",    "G2 V (V)",   "G2 I (A)",  # 7-10
    "G3 En", "G3 V (V)",   "G3 I (A)",              # 11-13
    "G4 En", "G4 V (V)",   "G4 I (A)",              # 14-16
    "GPP Tracking",                                   # 17
    "Notes",                                          # 18
    "ESPEC T (\u00b0C)",                              # 19
    "ESPEC H (%)",                                    # 20
]
_N     = len(_COLS)        # 21
_CHECK = {0, 3, 7, 11, 14}
_COMBO = {4, 8, 17}        # G1 Mode, G2 Mode, GPP Tracking
_PREC  = {1:3, 2:3, 5:3, 6:3, 9:3, 10:3, 12:3, 13:3, 15:3, 16:3, 19:1, 20:1}
_KEY   = {
     0:'itech_en',   1:'itech_v',    2:'itech_i',
     3:'g1_en',      4:'g1_mode',    5:'g1_v',       6:'g1_i',
     7:'g2_en',      8:'g2_mode',    9:'g2_v',      10:'g2_i',
    11:'g3_en',     12:'g3_v',      13:'g3_i',
    14:'g4_en',     15:'g4_v',      16:'g4_i',
    17:'gpp_tracking',
    18:'notes',
    19:'espec_t',
    20:'espec_h',
}
_DEFS  = {
     0:True,   1:0.0,    2:0.0,
     3:False,  4:'Source', 5:5.0,  6:1.0,
     7:False,  8:'Source', 9:12.0, 10:0.5,
    11:False, 12:3.3,   13:0.5,
    14:False, 15:15.0,  16:0.5,
    17:'Independent',
    18:'',
    19:'',
    20:'',
}
_COL_W = {
     0:44,  1:82,  2:82,
     3:44,  4:80,  5:74,  6:74,
     7:44,  8:80,  9:74, 10:74,
    11:44, 12:74, 13:74,
    14:44, 15:74, 16:74,
    17:100,
    19:80,
    20:80,
}
_TRACKING_OPTS = ["Independent", "Series", "Parallel"]
_MODE_OPTS = ["Source", "Load CV", "Load CC", "Load CR"]


class _ComboDelegate(QStyledItemDelegate):
    """Delegate for combo columns — items set per-column."""
    def __init__(self, options, parent=None):
        super().__init__(parent)
        self._options = options

    def createEditor(self, parent, option, index):
        cb = QComboBox(parent); cb.addItems(self._options); return cb
    def setEditorData(self, editor, index):
        editor.setCurrentIndex(max(0, editor.findText(index.data() or "")))
    def setModelData(self, editor, model, index):
        model.setData(index, editor.currentText())
    def updateEditorGeometry(self, editor, option, index):
        editor.setGeometry(option.rect)


_LOAD_BG = QColor('#3d2e1a')   # dark orange tint for load mode cells
_LOAD_FG = QColor('#fab387')   # orange text for load mode values


def _make_item(col: int, value=None) -> QTableWidgetItem:
    it = QTableWidgetItem()
    if col in _CHECK:
        it.setCheckState(Qt.CheckState.Checked if value else Qt.CheckState.Unchecked)
        it.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
    elif col in _COMBO:
        txt = str(value) if value is not None else _DEFS[col]
        it.setText(txt)
        it.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        # Color mode combo cells orange when not 'Source'
        if col in (4, 8) and txt != 'Source':
            it.setBackground(QBrush(_LOAD_BG))
            it.setForeground(QBrush(_LOAD_FG))
    elif col in _PREC:
        try:    it.setText(f"{float(value):.{_PREC[col]}f}")
        except: it.setText(str(_DEFS.get(col, '')))
        it.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
    else:
        it.setText(str(value) if value is not None else '')
    return it


def _read_row(tbl: QTableWidget, r: int) -> dict:
    def chk(c):
        it = tbl.item(r, c)
        return it is not None and it.checkState() == Qt.CheckState.Checked
    def num(c, d=0.0):
        it = tbl.item(r, c)
        try:    return float(it.text())
        except: return d
    def txt(c):
        it = tbl.item(r, c)
        return it.text().strip() if it else ''
    return {
        'itech_en':     chk(0),  'itech_v':      num(1),  'itech_i':    num(2),
        'g1_en':        chk(3),  'g1_mode':      txt(4),  'g1_v':       num(5),  'g1_i': num(6),
        'g2_en':        chk(7),  'g2_mode':      txt(8),  'g2_v':       num(9),  'g2_i': num(10),
        'g3_en':        chk(11), 'g3_v':         num(12), 'g3_i':       num(13),
        'g4_en':        chk(14), 'g4_v':         num(15), 'g4_i':       num(16),
        'gpp_tracking': txt(17),
        'notes':        txt(18),
        'espec_t':      txt(19),
        'espec_h':      txt(20),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Background runner
# ─────────────────────────────────────────────────────────────────────────────

class SequenceRunner(QThread):
    log_msg    = pyqtSignal(str)
    progress   = pyqtSignal(int, int)
    row_active = pyqtSignal(int)
    finished   = pyqtSignal(str)   # path to saved file (empty if none)

    def __init__(self, rows, defaults, timing, meas_inputs, eddie, out_path,
                 daq=None, espec=None, daq_en=False, espec_meas_t=False, espec_meas_h=False,
                 mdo_wf_en=False, mdo_wf_ch='CH1', mdo_wf_pts=10000,
                 gpp_panel=None):
        super().__init__()
        self._rows        = rows
        self._defaults    = defaults
        self._timing      = timing
        self._meas_inputs = meas_inputs
        self._eddie       = eddie
        self._out_path    = out_path
        self._stop        = threading.Event()
        self._daq          = daq
        self._espec        = espec
        self._daq_en       = daq_en
        self._espec_meas_t = espec_meas_t
        self._espec_meas_h = espec_meas_h
        self._mdo_wf_en    = mdo_wf_en
        self._mdo_wf_ch    = mdo_wf_ch
        self._mdo_wf_pts   = mdo_wf_pts
        self._gpp_panel    = gpp_panel

    def stop(self): self._stop.set()

    @staticmethod
    def _ts(): return datetime.now().strftime("[%H:%M:%S.%f")[:-3] + "]"

    def _emit(self, msg: str):
        """Emit log to UI and also to the rotating file logger."""
        self.log_msg.emit(msg)
        if '\u26a0' in msg or '\u274c' in msg or 'error' in msg.lower() or 'fail' in msg.lower():
            _log.warning("SeqRunner: %s", msg)
        else:
            _log.debug("SeqRunner: %s", msg)

    def _dwell(self, ms):
        s = ms / 1000.0; t0 = time.perf_counter()
        while not self._stop.is_set():
            if time.perf_counter() - t0 >= s: break
            time.sleep(min(0.05, s))

    def run(self):
        _log.info("SequenceRunner started, %d steps", len(self._rows))
        try:
            self._execute()
        except Exception as e:
            _log.error("SequenceRunner crashed: %s", e, exc_info=True)
            self._emit(f"\u274c FATAL: {e}")
            self.finished.emit("")

    def _execute(self):
        total        = len(self._rows)
        n_rd         = int(self._timing.get('n_readings', 1))
        interval_s   = self._timing.get('interval_ms', 500) / 1000.0
        settling     = self._timing.get('settling_ms', 0)
        avg_readings = bool(self._timing.get('avg_readings', True))
        on_read_fail = self._timing.get('on_read_fail', 'continue')
        max_retries  = int(self._timing.get('max_retries', 3))

        # ── EDDIE key list ────────────────────────────────────────────────
        eddie_keys = []
        if self._eddie:
            try:
                eddie_keys = list(self._eddie.get_watchlist_snapshot().keys())
            except Exception:
                pass

        # ── File path (resolved once, before loop) ────────────────────────
        fp = self._out_path.strip() if self._out_path else ""
        if not fp:
            fp = str(APP_DIR / "results" /
                     f"seq_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        if not fp.endswith(".xlsx"):
            fp += ".xlsx"
        Path(fp).parent.mkdir(parents=True, exist_ok=True)
        fp = unique_path(fp)

        def _flush():
            try:
                wb.save(fp)
            except Exception as _e:
                self._emit(f"  ⚠ Mid-run save error: {_e}")

        # ── Excel setup ───────────────────────────────────────────────────
        mdo_wf_en  = self._mdo_wf_en
        mdo_wf_ch  = self._mdo_wf_ch
        mdo_wf_pts = self._mdo_wf_pts
        _wf_tmpfiles: list = []

        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sequence Data"
        ws_wf = None
        _rd_hdr = [] if avg_readings else ["Reading"]
        headers = ["Timestamp", "Step"] + _rd_hdr + [
                   "Notes",
                   "ITECH V_set", "ITECH I_set", "ITECH Mode",
                   "G1 Mode", "G1 V_set", "G1 I_set",
                   "G2 Mode", "G2 V_set", "G2 I_set",
                   "G3 V_set", "G3 I_set", "G4 V_set", "G4 I_set",
                   "GPP Tracking",
                   "ESPEC T_set (\u00b0C)", "ESPEC H_set (%)"]
        for lbl, _ in self._meas_inputs:
            headers.append(lbl)
        for k in eddie_keys:
            headers.append(f"EDDIE {k}")
        daq_cfg_seq = self._daq.get_scan_config() if (self._daq_en and self._daq) else []
        for c in daq_cfg_seq:
            headers.append(c["label"])
        if self._espec_meas_t and self._espec:
            headers.append("ESPEC T_actual (\u00b0C)")
        if self._espec_meas_h and self._espec:
            headers.append("ESPEC H_actual (%)")
        ws.append(headers)
        for cell in ws[1]:
            cell.font = XFont(bold=True)
            cell.fill = PatternFill("solid", fgColor="313244")
        _flush()
        self._emit(f"✏️  Logging to: {fp}")

        gpp_keys = [('g1_en','g1_v','g1_i'), ('g2_en','g2_v','g2_i'),
                    ('g3_en','g3_v','g3_i'), ('g4_en','g4_v','g4_i')]

        # ── Apply GPP OVP/OCP once before step loop (from Instruments tab) ──
        _prot_cfg = (self._gpp_panel.get_protection_config()
                     if self._gpp_panel is not None else {})
        if _prot_cfg:
            try:
                def _apply_prot(res, _cfg=_prot_cfg):
                    d = GWGPP4323(res)
                    for ch in range(1, 5):
                        cd = _cfg.get(ch, {})
                        if cd.get('ovp_en', False):
                            d.set_ovp(ch, cd['ovp']); d.enable_ovp(ch, True)
                        if cd.get('ocp_en', False):
                            d.set_ocp(ch, cd['ocp']); d.enable_ocp(ch, True)
                shared_resource('gpp').run(_apply_prot)
                self._emit("GPP OVP/OCP applied (from Instruments tab)")
            except Exception as e:
                self._emit(f"  ⚠ GPP protection setup: {e}")

        for idx, row in enumerate(self._rows):
            if self._stop.is_set():
                self._emit("⏹  Stopped by user."); break

            self.progress.emit(idx, total)
            self.row_active.emit(idx)
            note = f"  — {row['notes']}" if row.get('notes') else ""
            self._emit(f"{self._ts()}  Step {idx+1}/{total}{note}")

            gpp_tracking = row.get('gpp_tracking', 'Independent').lower()

            # ── ITECH ─────────────────────────────────────────────────────
            if row.get('itech_en'):
                try:
                    v = float(row['itech_v']); i = float(row['itech_i'])
                    mode = "SOURCE" if i >= 0 else "SINK"
                    signed_i = i if mode == "SOURCE" else -abs(i)
                    _conn = shared_driver('itech', lambda: ITechFixed(ADDR['itech']))
                    def _apply(drv, _v=v, _i=signed_i):
                        drv.set_volts(_v); drv.set_amps(_i); drv.on()
                    _conn.run(_apply)
                    self._emit(f"  ITECH ON   V={v:.3f} V  I={i:+.3f} A  ({mode})")
                except Exception as e:
                    self._emit(f"  ⚠ ITECH: {e}")
            else:
                try:
                    _conn = shared_driver('itech', lambda: ITechFixed(ADDR['itech']))
                    _conn.run(lambda drv: drv.off())
                    self._emit(f"  ITECH OFF")
                except Exception as e:
                    self._emit(f"  ⚠ ITECH (off): {e}")

            # ── GPP tracking ───────────────────────────────────────────────
            if gpp_tracking != 'independent':
                try:
                    shared_resource('gpp').run(
                        lambda res, _t=gpp_tracking: GWGPP4323(res).set_tracking_mode(_t))
                    self._emit(f"  GPP tracking: {gpp_tracking}")
                except Exception as e:
                    self._emit(f"  ⚠ GPP tracking: {e}")

            # ── GPP channels (batched) ─────────────────────────────────────
            _mode_map = {'Load CV': 'CV', 'Load CC': 'CC', 'Load CR': 'CR'}

            try:
                _gpp_conn = shared_resource('gpp')
                def _gpp_batch(res, _row=row, _track=gpp_tracking):
                    drv = GWGPP4323(res)
                    msgs = []
                    for ch, (ek, vk, ik) in enumerate(gpp_keys, start=1):
                        if ch == 2 and _track != 'independent':
                            continue
                        mode_str = _row.get(f'g{ch}_mode', 'Source') if ch in (1, 2) else 'Source'
                        scpi_mode = _mode_map.get(mode_str)
                        if _row.get(ek):
                            v = float(_row[vk]); i = float(_row[ik])
                            if scpi_mode:
                                drv.set_load_mode(ch, scpi_mode)
                                if scpi_mode == 'CR':
                                    drv.set_load_resistance(ch, v)
                                else:
                                    drv.set_voltage(ch, v)
                                drv.set_current(ch, i)
                                drv.enable_channel(ch, True)
                                msgs.append(
                                    f"  GPP CH{ch} {mode_str}  V={v:.3f}  I={i:.3f}")
                            else:
                                if ch in (1, 2):
                                    drv.set_load_mode(ch, 'OFF')
                                drv.set_channel(ch, v, i)
                                drv.enable_channel(ch, True)
                                msgs.append(
                                    f"  GPP CH{ch} Source ON  V={v:.3f} V  I={i:.3f} A")
                        else:
                            if ch in (1, 2):
                                drv.set_load_mode(ch, 'OFF')
                            drv.enable_channel(ch, False)
                            msgs.append(f"  GPP CH{ch} OFF")
                    return msgs
                msgs = _gpp_conn.run(_gpp_batch)
                if msgs:
                    for m in msgs:
                        self._emit(m)
            except Exception as e:
                self._emit(f"  \u26a0 GPP: {e}")

            # \u2500\u2500 ESPEC setpoints for this step \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
            espec_t_str = row.get('espec_t', '').strip()
            espec_h_str = row.get('espec_h', '').strip()
            if self._espec and (espec_t_str or espec_h_str):
                et = eh = None
                try:
                    with use_instrument("espec") as res:
                        _espec_setup(res)
                        if espec_t_str:
                            et = float(espec_t_str)
                            espec_set_temp(res, et)
                            self._emit(f"  ESPEC T \u2192 {et:.1f}\u00b0C")
                        if espec_h_str:
                            eh = float(espec_h_str)
                            espec_set_humi(res, eh)
                            self._emit(f"  ESPEC H \u2192 {eh:.1f}%")
                except Exception as e:
                    self._emit(f"  \u26a0 ESPEC setpoint: {e}")
                    et = eh = None
                stab = self._espec.get_stabilization_config()
                if stab.get('wait_stable', False) and (et is not None or eh is not None):
                    tol_t = stab['tol_t']; tol_h = stab['tol_h']
                    tout  = stab['timeout_s']; poll = stab.get('stable_poll_s', 5)
                    t0_s  = time.perf_counter()
                    self._emit(f"  \u23f3 Stabilizing (\u00b1{tol_t}\u00b0C / \u00b1{tol_h}%  {tout}s)\u2026")
                    while not self._stop.is_set() and (time.perf_counter() - t0_s < tout):
                        try:
                            ok_t = ok_h = True
                            if et is not None:
                                at = self._espec.read_actual_temp_for_runner()
                                ok_t = abs(at - et) <= tol_t
                            if eh is not None:
                                ah = self._espec.read_actual_humi_for_runner()
                                ok_h = abs(ah - eh) <= tol_h
                            if ok_t and ok_h:
                                self._emit("  \u2713 Stable"); break
                        except Exception as e:
                            self._emit(f"  \u26a0 Stab: {e}")
                        time.sleep(poll)

            # ── Settling ───────────────────────────────────────────────────
            if settling > 0 and not self._stop.is_set():
                self._dwell(settling)
                self._emit(f"  (settled {settling} ms)")

            # ── N readings — accumulate, then average ─────────────────────
            itech_v_s  = row.get('itech_v', 0.0) if row.get('itech_en') else ''
            itech_i_s  = row.get('itech_i', 0.0) if row.get('itech_en') else ''
            itech_mode = ("SOURCE" if float(itech_i_s) >= 0 else "SINK") if itech_v_s != '' else ''

            # \u2500\u2500 DAQ scan (once per step) + ESPEC actual \u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500
            daq_vals = {}
            if self._daq_en and self._daq:
                try:
                    step_cfg = self._daq.get_scan_config()
                    if step_cfg:
                        with use_instrument("daq") as res:
                            daq_vals = daq_scan(res, step_cfg)
                        self._emit("  DAQ: " +
                                          "  ".join(f"{c['label']}={daq_vals.get(c['label'],float('nan')):.5g}"
                                                    for c in step_cfg))
                except Exception as e:
                    self._emit(f"  \u26a0 DAQ: {e}")
            espec_t_act = espec_h_act = float("nan")
            if self._espec_meas_t and self._espec:
                try:    espec_t_act = self._espec.read_actual_temp_for_runner()
                except Exception as e: self._emit(f"  \u26a0 ESPEC T_act: {e}")
            if self._espec_meas_h and self._espec:
                try:    espec_h_act = self._espec.read_actual_humi_for_runner()
                except Exception as e: self._emit(f"  \u26a0 ESPEC H_act: {e}")

            def _xrow_base(ts_val, rd_num=None):
                r = [ts_val, idx+1]
                if not avg_readings: r.append(rd_num)
                r += [row.get('notes',''), itech_v_s, itech_i_s, itech_mode]
                for ch in (1, 2):
                    r.append(row.get(f'g{ch}_mode', 'Source'))
                    r.append(row.get(f'g{ch}_v','') if row.get(f'g{ch}_en') else '')
                    r.append(row.get(f'g{ch}_i','') if row.get(f'g{ch}_en') else '')
                for ch in (3, 4):
                    r.append(row.get(f'g{ch}_v','') if row.get(f'g{ch}_en') else '')
                    r.append(row.get(f'g{ch}_i','') if row.get(f'g{ch}_en') else '')
                r.append(row.get('gpp_tracking',''))
                r.append(row.get('espec_t',''))
                r.append(row.get('espec_h',''))
                return r

            acc       = {lbl: [] for lbl, _ in self._meas_inputs}
            eddie_acc = {k:   [] for k in eddie_keys}
            _rd_buf   = []

            for _rd in range(n_rd):
                if self._stop.is_set(): break
                _rd_retry = 0
                while True:   # ── per-reading retry loop ───────────────────
                    t0 = time.perf_counter()
                    meas_now = {}; eddie_now = {}; _rd_fail = False

                    for lbl, fn in self._meas_inputs:
                        try:
                            v = fn()
                            if isinstance(v, float) and not math.isfinite(v):
                                _rd_fail = True
                        except Exception as e:
                            v = float('nan'); _rd_fail = True
                            self._emit(f"  ⚠ [{lbl}]: {e}")
                        meas_now[lbl] = v

                    if self._eddie and eddie_keys:
                        try:
                            snap = self._eddie.get_watchlist_snapshot()
                            for k in eddie_keys:
                                try:
                                    eddie_now[k] = float(snap.get(k, float('nan')))
                                except:
                                    eddie_now[k] = float('nan')
                        except Exception:
                            pass

                    # ── per-reading decision ──────────────────────────────
                    _do_acc        = True
                    _rd_loop_again = False
                    if _rd_fail and on_read_fail == 'skip':
                        self._emit(f"  ⏭  Rd {_rd+1} skipped (bad reading)")
                        _do_acc = False
                    elif _rd_fail and on_read_fail == 'redo':
                        if _rd_retry < max_retries:
                            _rd_retry += 1
                            self._emit(
                                f"  ⟳ Rd {_rd+1} retry ({_rd_retry}/{max_retries})…")
                            _rd_loop_again = True
                        else:
                            self._emit(
                                f"  ⚠ Rd {_rd+1} max retries ({max_retries}) — skipping")
                            _do_acc = False

                    if _do_acc:
                        for lbl, _ in self._meas_inputs:
                            acc[lbl].append(meas_now[lbl])
                        for k in eddie_keys:
                            eddie_acc[k].append(eddie_now.get(k, float('nan')))

                        if not avg_readings:
                            ts_rd = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                            xrow  = _xrow_base(ts_rd, _rd + 1)
                            for lbl, _ in self._meas_inputs:
                                xrow.append(meas_now.get(lbl, float('nan')))
                            for k in eddie_keys:
                                xrow.append(eddie_now.get(k, float('nan')))
                            for c in daq_cfg_seq:
                                xrow.append(daq_vals.get(c["label"], float("nan")))
                            if self._espec_meas_t and self._espec:
                                xrow.append(espec_t_act)
                            if self._espec_meas_h and self._espec:
                                xrow.append(espec_h_act)
                            _rd_buf.append(xrow)

                    rem = interval_s - (time.perf_counter() - t0)
                    if rem > 0 and not self._stop.is_set():
                        time.sleep(rem)

                    if not _rd_loop_again:
                        break

            # ── write data (step always logged) ───────────────────────────
            _n_good = len(next(iter(acc.values()), []))
            if avg_readings:
                def _avg(vals):
                    good = [v for v in vals if math.isfinite(v)]
                    return sum(good) / len(good) if good else float('nan')
                avg       = {lbl: _avg(v) for lbl, v in acc.items()}
                eddie_avg = {k:   _avg(v) for k,   v in eddie_acc.items()}
                ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                xrow = _xrow_base(ts)
                for lbl, _ in self._meas_inputs:
                    xrow.append(avg.get(lbl, float('nan')))
                for k in eddie_keys:
                    xrow.append(eddie_avg.get(k, float('nan')))
                for c in daq_cfg_seq:
                    xrow.append(daq_vals.get(c["label"], float("nan")))
                if self._espec_meas_t and self._espec:
                    xrow.append(espec_t_act)
                if self._espec_meas_h and self._espec:
                    xrow.append(espec_h_act)
                ws.append(xrow)
                if mdo_wf_en:
                    _png = _mdo_waveform_png(
                        mdo_wf_ch, mdo_wf_pts, f"Step {idx+1} avg")
                    if _png:
                        _wf_tmpfiles.append(_png)
                        try:
                            from openpyxl.drawing.image import Image as OpxImage
                            if ws_wf is None:
                                ws_wf = wb.create_sheet("Waveforms")
                                ws_wf.append(["Step", "Timestamp", "Ch"])
                            ws_wf.append([idx + 1,
                                          datetime.now().strftime("%H:%M:%S.%f")[:-3],
                                          mdo_wf_ch])
                            _ir = ws_wf.max_row + 1
                            ws_wf.row_dimensions[_ir].height = 150
                            _oi = OpxImage(_png)
                            _oi.width = 500; _oi.height = 200
                            ws_wf.add_image(_oi, f"A{_ir}")
                            ws_wf.row_dimensions[_ir + 1].height = 10
                        except Exception as _ie:
                            self._emit(f"  ⚠ Waveform embed: {_ie}")
                _flush()
                if avg:
                    _skipped = n_rd - _n_good
                    _sfx = f", {_skipped} skipped" if _skipped else ""
                    meas_str = "  ".join(f"{l}={v:.4g}" for l, v in avg.items())
                    self._emit(f"  ✓ {meas_str}  (n={_n_good}{_sfx})")
            else:
                for _ri, _xr in enumerate(_rd_buf):
                    ws.append(_xr)
                    if mdo_wf_en:
                        _png = _mdo_waveform_png(
                            mdo_wf_ch, mdo_wf_pts, f"Step {idx+1} Rd {_ri+1}")
                        if _png:
                            _wf_tmpfiles.append(_png)
                            try:
                                from openpyxl.drawing.image import Image as OpxImage
                                if ws_wf is None:
                                    ws_wf = wb.create_sheet("Waveforms")
                                    ws_wf.append(["Step", "Reading", "Timestamp", "Ch"])
                                ws_wf.append([idx + 1, _ri + 1,
                                              datetime.now().strftime("%H:%M:%S.%f")[:-3],
                                              mdo_wf_ch])
                                _ir = ws_wf.max_row + 1
                                ws_wf.row_dimensions[_ir].height = 150
                                _oi = OpxImage(_png)
                                _oi.width = 500; _oi.height = 200
                                ws_wf.add_image(_oi, f"A{_ir}")
                                ws_wf.row_dimensions[_ir + 1].height = 10
                            except Exception as _ie:
                                self._emit(f"  ⚠ Waveform embed: {_ie}")
                _flush()

        self.progress.emit(total, total)
        self.row_active.emit(-1)

        # ── Post-test state: apply first row values ─────────────────────
        self._emit("Applying post-test state (first row)…")
        first = self._rows[0] if self._rows else {}
        try:
            dv = float(first.get('itech_v', 0.0))
            di = float(first.get('itech_i', 0.0))
            en = bool(first.get('itech_en', False))
            _conn = shared_driver('itech', lambda: ITechFixed(ADDR['itech']))
            def _post(drv, _v=dv, _i=di, _en=en):
                drv.set_volts(_v); drv.set_amps(_i)
                if _en: drv.on()
                else:   drv.off()
            _conn.run(_post)
            self._emit(f"  ✓ ITECH {'ON' if en else 'OFF'}  V={dv:.3f} I={di:.3f}")
        except Exception as e:
            self._emit(f"  ⚠ ITECH post-test: {e}")
        try:
            _mode_map = {'Load CV': 'CV', 'Load CC': 'CC', 'Load CR': 'CR'}
            _tracking = first.get('gpp_tracking', 'Independent').lower()
            def _gpp_post(res, _first=first, _track=_tracking):
                drv = GWGPP4323(res)
                for ch in range(1, 5):
                    dv = float(_first.get(f'g{ch}_v', 0.0))
                    di = float(_first.get(f'g{ch}_i', 0.5))
                    en = bool(_first.get(f'g{ch}_en', False))
                    mode_str = _first.get(f'g{ch}_mode', 'Source') if ch in (1, 2) else 'Source'
                    scpi_mode = _mode_map.get(mode_str)
                    if scpi_mode:
                        drv.set_load_mode(ch, scpi_mode)
                        if scpi_mode == 'CR':
                            drv.set_load_resistance(ch, dv)
                        else:
                            drv.set_voltage(ch, dv)
                        drv.set_current(ch, di)
                    else:
                        if ch in (1, 2):
                            drv.set_load_mode(ch, 'OFF')
                        drv.set_channel(ch, dv, di)
                    drv.enable_channel(ch, en)
                drv.set_tracking_mode(_track)
            shared_resource('gpp').run(_gpp_post)
            self._emit(f"  \u2713 GPP post-test \u2192 first row (tracking={_tracking})")
        except Exception as e:
            self._emit(f"  \u26a0 GPP post-test: {e}")

        # ── Final save ───────────────────────────────────────────────────
        try:
            wb.save(fp)
            self._emit(f"✓  File closed: {fp}")
        except Exception as e:
            self._emit(f"⚠  Final save error: {e}")
        finally:
            import os
            for _tmp in _wf_tmpfiles:
                try: os.unlink(_tmp)
                except Exception: pass

        _log.info("SequenceRunner completed, output: %s", fp)
        self.finished.emit(fp)


# ─────────────────────────────────────────────────────────────────────────────
# Panel
# ─────────────────────────────────────────────────────────────────────────────

class SequencePanel(QWidget):
    test_running = pyqtSignal(bool)   # True when sequence starts, False when it ends

    def __init__(self, parent=None):
        super().__init__(parent)
        self._runner      = None
        self._eddie       = None
        self._daq_panel   = None
        self._espec_panel = None
        self._gpp_panel   = None
        self._lib_panel   = None
        self._last_file   = ""
        self._build_ui()

    def set_eddie_panel(self, panel):
        self._eddie = panel

    def set_daq_panel(self, panel):
        self._daq_panel = panel

    def set_espec_panel(self, panel):
        self._espec_panel = panel

    def set_gpp_panel(self, panel):
        """Called by MainWindow to pass the GppPanel reference for OVP/OCP config."""
        self._gpp_panel = panel

    def set_library_panel(self, panel):
        """Called by MainWindow to pass the InstrumentLibraryPanel reference."""
        self._lib_panel = panel

    # ── UI ────────────────────────────────────────────────────────────────────

    def _build_ui(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(6, 6, 6, 6)
        outer.setSpacing(5)

        # ── Toolbar ───────────────────────────────────────────────────────
        tb = QHBoxLayout(); tb.setSpacing(4)

        def _b(lbl, role=None, tip=""):
            b = QPushButton(lbl)
            if role: b.setProperty("role", role)
            if tip:  b.setToolTip(tip)
            b.setFixedHeight(28); return b

        def _sep():
            s = QLabel("|"); s.setStyleSheet("color:#45475a; padding:0 2px;"); return s

        self._btn_add   = _b("＋ Add",      tip="Append a blank step")
        self._btn_del   = _b("✕ Delete",    tip="Remove selected rows")
        self._btn_dup   = _b("⧉ Duplicate", tip="Clone selected rows below")
        self._btn_up    = _b("▲", tip="Move row up");   self._btn_up.setFixedWidth(30)
        self._btn_dn    = _b("▼", tip="Move row down"); self._btn_dn.setFixedWidth(30)
        self._btn_clr   = _b("🗑 Clear All", tip="Remove all rows")
        self._btn_run   = _b("▶  Run",  "green", tip="Execute sequence")
        self._btn_stop  = _b("■  Stop", "red",   tip="Abort run")
        self._btn_exp   = _b("📤 Export", tip="Save table as .xlsx")
        self._btn_imp   = _b("📥 Import", tip="Load table from .xlsx")

        for w in [self._btn_add, self._btn_del, self._btn_dup,
                  self._btn_up, self._btn_dn, self._btn_clr,
                  _sep(), self._btn_run, self._btn_stop,
                  _sep(), self._btn_exp, self._btn_imp]:
            tb.addWidget(w)
        tb.addStretch()
        self._status_lbl = QLabel("Ready")
        self._status_lbl.setStyleSheet("color:#a6adc8; font-style:italic;")
        tb.addWidget(self._status_lbl)

        # ── Progress ──────────────────────────────────────────────────────
        self._prog = QProgressBar()
        self._prog.setRange(0, 1); self._prog.setValue(0)
        self._prog.setFixedHeight(12); self._prog.setTextVisible(False)

        # ── Config row ────────────────────────────────────────────────────
        cfg = QHBoxLayout(); cfg.setSpacing(6)
        cfg.addWidget(self._build_timing_group(), 1)
        cfg.addWidget(self._build_meas_group(),   2)

        # ── Defaults strip ────────────────────────────────────────────────
        # (defined below)

        # ── Table ─────────────────────────────────────────────────────────
        self._tbl = QTableWidget(0, _N)
        self._tbl.setHorizontalHeaderLabels(_COLS)
        hdr = self._tbl.horizontalHeader()
        for c in range(_N):
            hdr.setSectionResizeMode(c, QHeaderView.ResizeMode.Interactive)
        hdr.setStretchLastSection(True)
        for c, w in _COL_W.items():
            self._tbl.setColumnWidth(c, w)
        self._tbl.verticalHeader().setDefaultSectionSize(26)
        self._tbl.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self._tbl.setAlternatingRowColors(True)
        self._tbl.setStyleSheet("QTableWidget { alternate-background-color: #1e1e2e; }")
        self._tbl.setItemDelegateForColumn(4,  _ComboDelegate(_MODE_OPTS, self._tbl))
        self._tbl.setItemDelegateForColumn(8,  _ComboDelegate(_MODE_OPTS, self._tbl))
        self._tbl.setItemDelegateForColumn(17, _ComboDelegate(_TRACKING_OPTS, self._tbl))
        self._tbl.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self._tbl.customContextMenuRequested.connect(self._ctx_menu)
        self._tbl.model().rowsInserted.connect(self._update_timing)
        self._tbl.model().rowsRemoved.connect(self._update_timing)

        # ── Log ───────────────────────────────────────────────────────────
        self._log_view = QTextEdit()
        self._log_view.setReadOnly(True)
        self._log_view.setFont(QFont("Courier New", 9))
        self._log_view.setMinimumHeight(80)

        # ── Wire ──────────────────────────────────────────────────────────
        self._btn_add.clicked.connect(lambda: self._add_row())
        self._btn_del.clicked.connect(self._del_rows)
        self._btn_dup.clicked.connect(self._dup_rows)
        self._btn_up.clicked.connect(self._move_up)
        self._btn_dn.clicked.connect(self._move_down)
        self._btn_clr.clicked.connect(self._clear)
        self._btn_run.clicked.connect(self._run)
        self._btn_stop.clicked.connect(self._stop_run)
        self._btn_stop.setEnabled(False)
        self._btn_exp.clicked.connect(self._export_excel)
        self._btn_imp.clicked.connect(self._import_excel)

        outer.addLayout(tb)
        outer.addWidget(self._prog)
        outer.addLayout(cfg)
        outer.addWidget(self._build_file_strip())
        outer.addWidget(self._build_col_toggles())

        # Table + log in a vertical splitter so user can resize
        tbl_log_split = QSplitter(Qt.Orientation.Vertical)
        tbl_log_split.setChildrenCollapsible(False)
        tbl_log_split.addWidget(self._tbl)

        log_w = QWidget()
        log_lay = QVBoxLayout(log_w)
        log_lay.setContentsMargins(0, 0, 0, 0)
        log_lay.setSpacing(2)
        log_lbl = QLabel("Run log:")
        log_lbl.setStyleSheet("color:#89b4fa; font-weight:bold;")
        log_lay.addWidget(log_lbl)
        log_lay.addWidget(self._log_view)
        tbl_log_split.addWidget(log_w)

        tbl_log_split.setStretchFactor(0, 3)
        tbl_log_split.setStretchFactor(1, 1)
        outer.addWidget(tbl_log_split, 1)

    # ── Column group toggles ─────────────────────────────────────────────────

    _COL_GROUPS = {
        'ITECH':      (0, 1, 2),
        'GPP CH1':    (3, 4, 5, 6),
        'GPP CH2':    (7, 8, 9, 10),
        'GPP CH3':    (11, 12, 13),
        'GPP CH4':    (14, 15, 16),
        'Tracking':   (17,),
        'ESPEC':      (19, 20),
    }

    def _build_col_toggles(self) -> QFrame:
        frm = QFrame()
        lay = QHBoxLayout(frm)
        lay.setContentsMargins(4, 2, 4, 2)
        lay.setSpacing(6)

        lbl = QLabel("Show columns:")
        lbl.setStyleSheet("color:#a6adc8; font-size:9pt; font-style:italic;")
        lay.addWidget(lbl)

        self._col_toggles = {}
        for name, cols in self._COL_GROUPS.items():
            cb = QCheckBox(name)
            cb.setChecked(True)
            cb.setStyleSheet("font-size:9pt;")
            cb.toggled.connect(lambda checked, cs=cols: self._toggle_cols(cs, checked))
            lay.addWidget(cb)
            self._col_toggles[name] = cb

        lay.addStretch()

        # Quick toggle all
        btn_all = QPushButton("All")
        btn_all.setFixedWidth(40); btn_all.setFixedHeight(22)
        btn_all.setStyleSheet("font-size:8pt; min-width:0; padding:2px 6px;")
        btn_all.clicked.connect(lambda: self._set_all_col_toggles(True))
        btn_min = QPushButton("Min")
        btn_min.setFixedWidth(40); btn_min.setFixedHeight(22)
        btn_min.setStyleSheet("font-size:8pt; min-width:0; padding:2px 6px;")
        btn_min.clicked.connect(lambda: self._set_all_col_toggles(False))
        lay.addWidget(btn_all)
        lay.addWidget(btn_min)
        return frm

    def _toggle_cols(self, cols, visible):
        for c in cols:
            if visible:
                self._tbl.showColumn(c)
            else:
                self._tbl.hideColumn(c)

    def _set_all_col_toggles(self, show: bool):
        for cb in self._col_toggles.values():
            cb.setChecked(show)

    # ── Timing group ──────────────────────────────────────────────────────────

    def _build_timing_group(self) -> QGroupBox:
        grp = QGroupBox("Sequence Timing")
        g = QGridLayout(grp); g.setContentsMargins(6, 4, 6, 4)

        g.addWidget(QLabel("Readings/step:"), 0, 0)
        self.n_readings = QSpinBox()
        self.n_readings.setRange(1, 1000); self.n_readings.setValue(3)
        self.n_readings.valueChanged.connect(self._update_timing)
        g.addWidget(self.n_readings, 0, 1)

        g.addWidget(QLabel("Interval:"), 0, 2)
        self.interval_ms = QSpinBox()
        self.interval_ms.setRange(100, 60000); self.interval_ms.setValue(1000)
        self.interval_ms.setSuffix(" ms")
        self.interval_ms.valueChanged.connect(self._update_timing)
        g.addWidget(self.interval_ms, 0, 3)

        g.addWidget(QLabel("Settling:"), 0, 4)
        self.settling_ms = QSpinBox()
        self.settling_ms.setRange(0, 60000); self.settling_ms.setValue(500)
        self.settling_ms.setSuffix(" ms")
        self.settling_ms.setToolTip("Wait after setting sources, before readings start")
        self.settling_ms.valueChanged.connect(self._update_timing)
        g.addWidget(self.settling_ms, 0, 5)

        g.addWidget(QLabel("Total time:"), 1, 0)
        self._timing_lbl = QLabel("—")
        self._timing_lbl.setStyleSheet("color:#fab387; font-weight:bold;")
        g.addWidget(self._timing_lbl, 1, 1, 1, 3)

        self.avg_readings = QCheckBox("Average readings  (one Excel row per step)")
        self.avg_readings.setChecked(True)
        self.avg_readings.setToolTip(
            "If unchecked, every individual reading is written as its own row in Excel")
        g.addWidget(self.avg_readings, 2, 0, 1, 6)

        g.addWidget(QLabel("On read fail:"), 3, 0)
        self.on_read_fail = QComboBox()
        self.on_read_fail.addItems(["Continue  (use partial)", "Skip reading", "Redo reading"])
        self.on_read_fail.setToolTip(
            "What to do when a single reading is bad (NaN / Inf):\n"
            "  Continue      — include all readings; bad values filtered out in average (default)\n"
            "  Skip reading  — discard that reading; step still logged with remaining good readings\n"
            "  Redo reading  — retry that individual reading up to Max retries times")
        g.addWidget(self.on_read_fail, 3, 1, 1, 2)
        self._retry_lbl  = QLabel("Max retries:")
        self.max_retries = QSpinBox()
        self.max_retries.setRange(1, 10); self.max_retries.setValue(3)
        self._retry_lbl.setVisible(False); self.max_retries.setVisible(False)
        self.on_read_fail.currentIndexChanged.connect(
            lambda i: (self._retry_lbl.setVisible(i == 2),
                       self.max_retries.setVisible(i == 2)))
        g.addWidget(self._retry_lbl,  3, 3)
        g.addWidget(self.max_retries, 3, 4)

        self._update_timing()
        return grp

    def _update_timing(self, *_):
        n_steps     = self._tbl.rowCount() if hasattr(self, '_tbl') else 0
        per_step_ms = (self.settling_ms.value() +
                       self.n_readings.value() * self.interval_ms.value())
        total_ms    = n_steps * per_step_ms
        total_s     = total_ms / 1000.0
        mins, secs  = divmod(int(total_s), 60)
        hrs,  mins  = divmod(mins, 60)
        t_str = (f"{hrs}h {mins:02d}m {secs:02d}s" if hrs
                 else f"{mins}m {secs:02d}s" if mins
                 else f"{secs}s")
        self._timing_lbl.setText(
            f"{n_steps} steps  ×  {per_step_ms} ms/step  =  {t_str}  "
            f"({total_s:.1f} s)")

    # ── Measurement group ─────────────────────────────────────────────────────

    def _build_meas_group(self) -> QGroupBox:
        grp = QGroupBox("Measurement Inputs  (averaged over N readings, one row per step)")
        g = QGridLayout(grp); g.setContentsMargins(6, 4, 6, 4); row = 0

        self.gdm_en = QCheckBox("GDM-8261A"); self.gdm_en.setChecked(True)
        self.gdm_func = QComboBox()
        self.gdm_func.addItems(["VOLT:DC","VOLT:AC","CURR:DC","CURR:AC","RES","FREQ"])
        g.addWidget(self.gdm_en, row, 0); g.addWidget(self.gdm_func, row, 1); row += 1

        self.dmm_en = QCheckBox("Keysight 34461A"); self.dmm_en.setChecked(True)
        self.dmm_func = QComboBox()
        self.dmm_func.addItems(["VOLT:DC","VOLT:AC","CURR:DC","CURR:AC","RES","FREQ"])
        g.addWidget(self.dmm_en, row, 0); g.addWidget(self.dmm_func, row, 1); row += 1

        g.addWidget(QLabel("ITECH:"), row, 0)
        self.itech_meas_v = QCheckBox("V"); self.itech_meas_v.setChecked(True)
        self.itech_meas_i = QCheckBox("I"); self.itech_meas_i.setChecked(True)
        g.addWidget(self.itech_meas_v, row, 1); g.addWidget(self.itech_meas_i, row, 2); row += 1

        g.addWidget(QLabel("GPP V:"), row, 0)
        self._gpp_meas_v = {}
        for ch in range(1, 5):
            cb = QCheckBox(f"CH{ch}"); cb.setChecked(ch == 1)
            self._gpp_meas_v[ch] = cb; g.addWidget(cb, row, ch)
        row += 1

        g.addWidget(QLabel("GPP I:"), row, 0)
        self._gpp_meas_i = {}
        for ch in range(1, 5):
            cb = QCheckBox(f"CH{ch}"); cb.setChecked(False)
            self._gpp_meas_i[ch] = cb; g.addWidget(cb, row, ch)
        row += 1

        self.seq_daq_en = QCheckBox("DAQ 973A  (all enabled channels)")
        self.seq_daq_en.setToolTip(
            "Scan all channels configured in the DAQ 973A instrument panel.\n"
            "One Excel column per enabled channel.")
        g.addWidget(self.seq_daq_en, row, 0, 1, 4)
        row += 1

        g.addWidget(QLabel("ESPEC actual:"), row, 0)
        self.seq_espec_meas_t = QCheckBox("T_actual (\u00b0C)")
        self.seq_espec_meas_h = QCheckBox("H_actual (%)")
        g.addWidget(self.seq_espec_meas_t, row, 1)
        g.addWidget(self.seq_espec_meas_h, row, 2)
        row += 1

        # Yokogawa WT5000
        self._wt_container = QFrame()
        wt_lay = QHBoxLayout(self._wt_container)
        wt_lay.setContentsMargins(0, 0, 0, 0); wt_lay.setSpacing(4)
        self.wt_en = QCheckBox("WT5000  E:")
        wt_lay.addWidget(self.wt_en)
        self.wt_elem = QSpinBox(); self.wt_elem.setRange(1, 7); self.wt_elem.setValue(1)
        self.wt_elem.setFixedWidth(132)
        wt_lay.addWidget(self.wt_elem)
        self._wt_checks = {}
        for _func, _lbl in [("URMS","Vrms"),("IRMS","Arms"),("UDC","Udc"),("IDC","Idc"),
                             ("P","W"),("LAMBda","PF"),("S","VA"),("Q","VAR")]:
            _cb = QCheckBox(_lbl)
            _cb.setChecked(_func in ("URMS", "IRMS", "P", "LAMBda"))
            self._wt_checks[_func] = _cb
            wt_lay.addWidget(_cb)
        wt_lay.addStretch()
        g.addWidget(self._wt_container, row, 0, 1, 4)
        row += 1

        # Tektronix MDO34
        self._mdo_container = QFrame()
        mdo_lay = QHBoxLayout(self._mdo_container)
        mdo_lay.setContentsMargins(0, 0, 0, 0); mdo_lay.setSpacing(4)
        self.mdo_en = QCheckBox("MDO34")
        mdo_lay.addWidget(self.mdo_en)
        mdo_lay.addWidget(QLabel("Ch:"))
        self.mdo_ch = QComboBox()
        self.mdo_ch.addItems(MDO34Driver.CHANNELS)
        self.mdo_ch.setFixedWidth(64)
        mdo_lay.addWidget(self.mdo_ch)
        mdo_lay.addWidget(QLabel("Meas:"))
        self.mdo_meas = QComboBox()
        for lbl, val in MDO34Driver.MEAS_TYPES:
            self.mdo_meas.addItem(lbl, val)
        mdo_lay.addWidget(self.mdo_meas)
        mdo_lay.addStretch()
        g.addWidget(self._mdo_container, row, 0, 1, 4)

        return grp

    # ── File group ────────────────────────────────────────────────────────────

    def _browse_file(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save sequence results", "", "Excel Workbook (*.xlsx)")
        if path:
            if not path.endswith(".xlsx"): path += ".xlsx"
            self.file_edit.setText(path)

    # ── Output file strip (replaces old defaults strip) ─────────────────────

    def _build_file_strip(self) -> QGroupBox:
        grp = QGroupBox("Output File  (.xlsx)  —  post-test state is set to the first row of the sequence")
        lay = QHBoxLayout(grp); lay.setContentsMargins(6, 4, 6, 4)
        lay.addWidget(QLabel("Save as:"))
        self.file_edit = QLineEdit()
        self.file_edit.setPlaceholderText("Leave blank → auto-generate in results/")
        lay.addWidget(self.file_edit, stretch=1)
        btn = QPushButton("Browse…"); btn.setFixedHeight(26)
        btn.clicked.connect(self._browse_file)
        lay.addWidget(btn)
        return grp

    def _get_defaults(self) -> dict:
        """Derive post-test defaults from the first row of the table.
        After the test completes, instruments are set to row 0 values."""
        if self._tbl.rowCount() == 0:
            return {
                'itech_v': 0, 'itech_i': 0, 'itech_en_after': False,
                'gpp_tracking': 'independent',
                'g1_v': 0, 'g1_i': 0, 'g1_en_after': False,
                'g2_v': 0, 'g2_i': 0, 'g2_en_after': False,
                'g3_v': 0, 'g3_i': 0, 'g3_en_after': False,
                'g4_v': 0, 'g4_i': 0, 'g4_en_after': False,
            }
        row = _read_row(self._tbl, 0)
        d = {
            'itech_v':        row.get('itech_v', 0),
            'itech_i':        row.get('itech_i', 0),
            'itech_en_after': row.get('itech_en', False),
            'gpp_tracking':   row.get('gpp_tracking', 'Independent').lower(),
        }
        for ch in range(1, 5):
            d[f'g{ch}_v']        = row.get(f'g{ch}_v', 0)
            d[f'g{ch}_i']        = row.get(f'g{ch}_i', 0)
            d[f'g{ch}_en_after'] = row.get(f'g{ch}_en', False)
        return d

    # ── Measurement callables ─────────────────────────────────────────────────

    def _build_meas_inputs(self) -> list:
        meas = []
        if self.gdm_en.isChecked():
            f = self.gdm_func.currentText()
            meas.append((f"GDM {f}",  lambda func=f: self._meas_gdm(func)))
        if self.dmm_en.isChecked():
            f = self.dmm_func.currentText()
            meas.append((f"DMM {f}",  lambda func=f: self._meas_dmm(func)))
        if self.itech_meas_v.isChecked():
            meas.append(("ITECH V",   self._meas_itech_v))
        if self.itech_meas_i.isChecked():
            meas.append(("ITECH I",   self._meas_itech_i))
        for ch in range(1, 5):
            if self._gpp_meas_v[ch].isChecked():
                meas.append((f"GPP CH{ch} V", lambda c=ch: self._meas_gpp_v(c)))
            if self._gpp_meas_i[ch].isChecked():
                meas.append((f"GPP CH{ch} I", lambda c=ch: self._meas_gpp_i(c)))
        if self.wt_en.isChecked():
            _elem  = self.wt_elem.value()
            _funcs = [f for f, cb in self._wt_checks.items() if cb.isChecked()]
            if _funcs:
                meas.extend(self._build_wt5000_meas(_elem, _funcs))
        if self.mdo_en.isChecked():
            _ch   = self.mdo_ch.currentText()
            _mtyp = self.mdo_meas.currentData()
            _lbl  = self.mdo_meas.currentText()
            meas.append(
                (f"MDO34 {_ch} {_lbl}", lambda c=_ch, m=_mtyp: self._meas_mdo(c, m)))
        return meas

    def _meas_gdm(self, func):
        needs_conf = getattr(self, '_gdm_last_func', None) != func
        self._gdm_last_func = func
        def fn(res, _f=func, _conf=needs_conf):
            if _conf:
                res.write(f"CONF:{_f}")
                time.sleep(0.1)
            return float(res.query("READ?").strip().split(",")[0])
        return shared_resource('gdm').run(fn)

    def _meas_dmm(self, func):
        with use_instrument("dmm") as res:
            return float(res.query(f"MEAS:{func}? AUTO,DEF").strip())

    def _meas_itech_v(self):
        conn = shared_driver('itech', lambda: ITechFixed(ADDR['itech']))
        result = conn.run(lambda drv: float(drv.ask('MEAS:VOLT?').strip()))
        return result if result is not None else float('nan')

    def _meas_itech_i(self):
        conn = shared_driver('itech', lambda: ITechFixed(ADDR['itech']))
        result = conn.run(lambda drv: float(drv.ask('MEAS:CURR?').strip()))
        return result if result is not None else float('nan')

    def _meas_gpp_v(self, ch):
        conn = shared_resource('gpp')
        result = conn.run(lambda res, c=ch: float(res.query(f"MEAS{c}:VOLT?").strip()))
        return result if result is not None else float('nan')

    def _meas_gpp_i(self, ch):
        conn = shared_resource('gpp')
        result = conn.run(lambda res, c=ch: float(res.query(f"MEAS{c}:CURR?").strip()))
        return result if result is not None else float('nan')

    def _meas_mdo(self, ch: str, meas_type: str) -> float:
        conn = shared_driver('mdo', lambda: MDO34Driver(ADDR['mdo']))
        result = conn.run(lambda drv: drv.measure_immediate(ch, meas_type))
        return result if result is not None else float('nan')

    def _build_wt5000_meas(self, elem: int, funcs: list) -> list:
        """Returns (label, fn) pairs for selected WT5000 functions.
        All share one persistent connection per run via shared_driver, which
        avoids TCP TIME_WAIT failures that occur when open/close is called for
        every reading.  The first function in the list fetches all values; the
        rest read from the cached state dict."""
        label_map = dict(WT5000Driver.ELEM_ITEMS)
        _state = {'data': {}}

        def _fetch():
            _state['data'] = {}   # clear so failures yield NaN, not stale data
            conn   = shared_driver('wt5000', lambda: WT5000Driver(ADDR['wt5000']))
            result = conn.run(lambda drv: drv.measure_element(elem, funcs))
            if result is not None:
                _state['data'] = result

        result = []
        for i, func in enumerate(funcs):
            full_lbl  = f"WT5000 E{elem} {label_map.get(func, func)}"
            cache_lbl = label_map.get(func, func)
            if i == 0:
                def _make_first(cl):
                    def fn():
                        _fetch()
                        return _state['data'].get(cl, float('nan'))
                    return fn
                result.append((full_lbl, _make_first(cache_lbl)))
            else:
                def _make_rest(cl):
                    return lambda: _state['data'].get(cl, float('nan'))
                result.append((full_lbl, _make_rest(cache_lbl)))
        return result

    # ── Session save / restore (table rows) ────────────────────────────────

    def save_steps_session(self):
        """Serialise all table rows to sequence_session.json."""
        rows = [_read_row(self._tbl, r) for r in range(self._tbl.rowCount())]
        try:
            _STEPS_SESSION.write_text(
                json.dumps(rows, indent=2, ensure_ascii=False), encoding='utf-8')
        except Exception as e:
            print(f"[session] save error: {e}")

    def restore_steps_session(self):
        """Populate the table from sequence_session.json (called at startup)."""
        if not _STEPS_SESSION.exists():
            return
        try:
            rows = json.loads(_STEPS_SESSION.read_text(encoding='utf-8'))
            self._tbl.setRowCount(0)
            for d in rows:
                self._add_row(d)
        except Exception as e:
            print(f"[session] load error: {e}")

    def connect_steps_autosave(self, callback):
        """Connect table change signals to *callback* so the autosave timer fires
        whenever rows are added, removed, or edited."""
        self._tbl.itemChanged.connect(lambda _: callback())
        self._tbl.model().rowsInserted.connect(lambda *_: callback())
        self._tbl.model().rowsRemoved.connect(lambda *_: callback())

    def notify_instrument_connected(self, key: str):
        pass  # all instruments always visible now

    def notify_instrument_disconnected(self, key: str):
        if key == 'wt5000':
            self.wt_en.setChecked(False)
        elif key == 'mdo':
            self.mdo_en.setChecked(False)

    # ── Context menu ──────────────────────────────────────────────────────────

    def _ctx_menu(self, pos):
        if not self._selected_rows(): return
        menu = QMenu(self)
        menu.addAction("↺  Reset to Defaults", self._reset_rows)
        menu.addAction("⧉  Duplicate",         self._dup_rows)
        menu.addSeparator()
        menu.addAction("✕  Delete Row(s)",      self._del_rows)
        menu.exec(self._tbl.viewport().mapToGlobal(pos))

    def _reset_rows(self):
        for r in self._selected_rows():
            for c in range(_N):
                self._tbl.setItem(r, c, _make_item(c, _DEFS[c]))

    # ── Row helpers ───────────────────────────────────────────────────────────

    def _add_row(self, d: dict = None):
        r = self._tbl.rowCount(); self._tbl.insertRow(r)
        for c in range(_N):
            self._tbl.setItem(r, c, _make_item(c, d.get(_KEY[c], _DEFS[c]) if d else _DEFS[c]))

    def _fill_row(self, r: int, d: dict):
        for c in range(_N):
            self._tbl.setItem(r, c, _make_item(c, d.get(_KEY[c], _DEFS[c])))

    def _selected_rows(self):
        return sorted({i.row() for i in self._tbl.selectedItems()})

    def _del_rows(self):
        for r in sorted(self._selected_rows(), reverse=True):
            self._tbl.removeRow(r)

    def _dup_rows(self):
        rows = self._selected_rows()
        ins  = (rows[-1] + 1) if rows else self._tbl.rowCount()
        for off, r in enumerate(rows):
            self._tbl.insertRow(ins + off)
            self._fill_row(ins + off, _read_row(self._tbl, r))

    def _move_up(self):
        rows = self._selected_rows()
        if not rows or rows[0] == 0: return
        for r in rows: self._swap(r - 1, r)
        self._tbl.clearSelection()
        for r in rows: self._tbl.selectRow(r - 1)

    def _move_down(self):
        rows = self._selected_rows()
        if not rows or rows[-1] == self._tbl.rowCount() - 1: return
        for r in reversed(rows): self._swap(r, r + 1)
        self._tbl.clearSelection()
        for r in rows: self._tbl.selectRow(r + 1)

    def _swap(self, a, b):
        da = _read_row(self._tbl, a); db = _read_row(self._tbl, b)
        self._fill_row(a, db); self._fill_row(b, da)

    def _clear(self):
        if self._tbl.rowCount() == 0: return
        if QMessageBox.question(self, "Clear All", "Remove all rows?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) == QMessageBox.StandardButton.Yes:
            self._tbl.setRowCount(0)

    def _highlight_row(self, r: int):
        active = QBrush(QColor("#2a3f5f")); clear = QBrush(QColor(0,0,0,0))
        for row in range(self._tbl.rowCount()):
            bg = active if row == r else clear
            for c in range(_N):
                it = self._tbl.item(row, c)
                if it: it.setBackground(bg)

    # ── Run / Stop ────────────────────────────────────────────────────────────

    def _run(self):
        n = self._tbl.rowCount()
        if n == 0:
            QMessageBox.warning(self, "Empty Sequence", "Add at least one step."); return

        # ── Pre-flight: check required instruments are configured ─────────
        rows_pre = [_read_row(self._tbl, r) for r in range(n)]
        required = []
        if any(r.get('itech_en') for r in rows_pre):
            required.append('itech')
        for ch in range(1, 5):
            if any(r.get(f'g{ch}_en') for r in rows_pre):
                if 'gpp' not in required:
                    required.append('gpp')
                break
        if any(r.get('espec_t', '').strip() or r.get('espec_h', '').strip()
               for r in rows_pre):
            required.append('espec')
        # Check measurement instruments
        meas_pre = self._build_meas_inputs()
        for lbl, _ in meas_pre:
            ll = lbl.lower()
            if 'gdm' in ll and 'gdm' not in required:     required.append('gdm')
            if 'dmm' in ll and 'dmm' not in required:     required.append('dmm')
            if 'mdo' in ll and 'mdo' not in required:     required.append('mdo')
            if 'wt5000' in ll and 'wt5000' not in required: required.append('wt5000')
        if self.seq_daq_en.isChecked():
            required.append('daq')
        # Check address configuration
        no_addr = check_instruments(required)
        if no_addr:
            names = "\n".join(f"  \u2022 {n}  ({k}) \u2014 no address" for k, n in no_addr)
            mb = QMessageBox(self)
            mb.setWindowTitle("Instruments Not Configured")
            mb.setIcon(QMessageBox.Icon.Critical)
            mb.setText(f"The following instruments have no valid address:\n\n{names}\n\n"
                       "Configure them in the Addresses tab first.")
            mb.addButton("Cancel", QMessageBox.ButtonRole.RejectRole)
            mb.exec()
            return

        # Check instruments are connected via Instruments tab
        connected = (self._lib_panel.connected_keys()
                     if self._lib_panel is not None else set())
        not_connected = [(k, _INSTR_NAMES.get(k, k))
                         for k in required if k not in connected]
        if not_connected:
            names = "\n".join(f"  \u2022 {n}  ({k})" for k, n in not_connected)
            mb = QMessageBox(self)
            mb.setWindowTitle("Instruments Not Connected")
            mb.setIcon(QMessageBox.Icon.Warning)
            mb.setText(f"The following instruments are not connected via the Instruments tab:\n\n"
                       f"{names}\n\n"
                       "Connect them first for reliable operation, or proceed anyway\n"
                       "(the test will attempt to open connections on the fly).")
            btn_proceed = mb.addButton("Proceed Anyway", QMessageBox.ButtonRole.AcceptRole)
            mb.addButton("Cancel", QMessageBox.ButtonRole.RejectRole)
            mb.exec()
            if mb.clickedButton() != btn_proceed:
                return

        if self._eddie is not None and getattr(self._eddie, '_pcu', None) is None:
            mb = QMessageBox(self)
            mb.setWindowTitle("EDDIE Not Connected")
            mb.setText("The EDDIE device is not connected.\n"
                       "EDDIE measurements will be skipped for this run.")
            mb.setIcon(QMessageBox.Icon.Warning)
            mb.addButton("Proceed Anyway", QMessageBox.ButtonRole.AcceptRole)
            mb.addButton("Cancel",         QMessageBox.ButtonRole.RejectRole)
            mb.exec()
            if mb.clickedButton().text() == "Cancel":
                return

        rows    = rows_pre
        timing  = {'n_readings':   self.n_readings.value(),
                   'interval_ms':  self.interval_ms.value(),
                   'settling_ms':  self.settling_ms.value(),
                   'avg_readings': self.avg_readings.isChecked(),
                   'on_read_fail': ['continue', 'skip', 'redo'][
                                       self.on_read_fail.currentIndex()],
                   'max_retries':  self.max_retries.value()}
        meas    = self._build_meas_inputs()
        out     = self.file_edit.text().strip()

        self._prog.setRange(0, n); self._prog.setValue(0)
        self._btn_run.setEnabled(False); self._btn_stop.setEnabled(True)
        self._status("Running…", "#a6e3a1")
        self._log_view.clear(); self._highlight_row(-1)

        _mdo_cfg = {}
        if hasattr(self, 'mdo_en') and self.mdo_en.isChecked():
            _mdo_cfg = {'enabled': True,
                        'ch':  self.mdo_ch.currentText(),
                        'pts': 10000}
        self._runner = SequenceRunner(rows, self._get_defaults(), timing, meas,
                                      self._eddie, out,
                                      daq=self._daq_panel,
                                      espec=self._espec_panel,
                                      daq_en=self.seq_daq_en.isChecked(),
                                      espec_meas_t=self.seq_espec_meas_t.isChecked(),
                                      espec_meas_h=self.seq_espec_meas_h.isChecked(),
                                      mdo_wf_en=_mdo_cfg.get('enabled', False),
                                      mdo_wf_ch=_mdo_cfg.get('ch', 'CH1'),
                                      mdo_wf_pts=_mdo_cfg.get('pts', 10000),
                                      gpp_panel=self._gpp_panel)
        self._runner.log_msg.connect(self._append_log)
        self._runner.progress.connect(lambda v, m: self._prog.setValue(v))
        self._runner.row_active.connect(self._highlight_row)
        self._runner.finished.connect(self._on_done)
        self._runner.start()
        self.test_running.emit(True)

    def _stop_run(self):
        if self._runner: self._runner.stop()

    def _on_done(self, path: str):
        self.test_running.emit(False)
        self._last_file = path
        self._btn_run.setEnabled(True); self._btn_stop.setEnabled(False)
        msg = f"✓ Complete  →  {path}" if path else "✓ Complete"
        self._status(msg, "#a6e3a1")

    def _append_log(self, msg: str):
        self._log_view.append(msg)
        self._log_view.moveCursor(QTextCursor.MoveOperation.End)

    def _status(self, text: str, color: str = "#a6adc8"):
        self._status_lbl.setText(text)
        self._status_lbl.setStyleSheet(f"color:{color}; font-style:italic;")

    # ── Excel export (table + config sheet) ──────────────────────────────────

    def _export_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Sequence", "", "Excel Files (*.xlsx)")
        if not path: return
        if not path.endswith(".xlsx"): path += ".xlsx"
        try:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sequence"
            ws.append(_COLS)
            for cell in ws[1]:
                cell.font = XFont(bold=True)
                cell.fill = PatternFill("solid", fgColor="313244")
            for r in range(self._tbl.rowCount()):
                row_data = []
                for c in range(_N):
                    it = self._tbl.item(r, c)
                    if c in _CHECK:
                        row_data.append(1 if (it and it.checkState() == Qt.CheckState.Checked) else 0)
                    else:
                        row_data.append(it.text() if it else "")
                ws.append(row_data)
            for col_cells in ws.columns:
                w = max(len(str(c.value or "")) for c in col_cells)
                ws.column_dimensions[col_cells[0].column_letter].width = max(8, w+2)

            # ── Config sheet (timing + measurement inputs + post-test defaults) ──
            wc = wb.create_sheet("Config")
            wc.append(["Key", "Value"])
            wc["A1"].font = XFont(bold=True); wc["B1"].font = XFont(bold=True)
            cfg_rows = [
                ["n_readings",      self.n_readings.value()],
                ["interval_ms",     self.interval_ms.value()],
                ["settling_ms",     self.settling_ms.value()],
                ["avg_readings",    int(self.avg_readings.isChecked())],
                ["gdm_en",          int(self.gdm_en.isChecked())],
                ["gdm_func",        self.gdm_func.currentText()],
                ["dmm_en",          int(self.dmm_en.isChecked())],
                ["dmm_func",        self.dmm_func.currentText()],
                ["itech_meas_v",    int(self.itech_meas_v.isChecked())],
                ["itech_meas_i",    int(self.itech_meas_i.isChecked())],
            ]
            for ch in range(1, 5):
                cfg_rows.append([f"gpp_meas_v_{ch}", int(self._gpp_meas_v[ch].isChecked())])
                cfg_rows.append([f"gpp_meas_i_{ch}", int(self._gpp_meas_i[ch].isChecked())])
            cfg_rows += [
                ["seq_daq_en",       int(self.seq_daq_en.isChecked())],
                ["seq_espec_meas_t", int(self.seq_espec_meas_t.isChecked())],
                ["seq_espec_meas_h", int(self.seq_espec_meas_h.isChecked())],
                ["wt_en",            int(self.wt_en.isChecked())],
                ["wt_elem",          self.wt_elem.value()],
            ]
            for func, cb in self._wt_checks.items():
                cfg_rows.append([f"wt_{func}", int(cb.isChecked())])
            if self._mdo_container.isVisible():
                cfg_rows += [
                    ["mdo_en",       int(self.mdo_en.isChecked())],
                    ["mdo_ch",       self.mdo_ch.currentText()],
                    ["mdo_meas_idx", self.mdo_meas.currentIndex()],
                ]
            cfg_rows.append(["on_read_fail_idx",   self.on_read_fail.currentIndex()])
            cfg_rows.append(["max_retries",         self.max_retries.value()])
            for row_kv in cfg_rows:
                wc.append(row_kv)

            wb.save(path)
            self._status(f"Exported → {Path(path).name}", "#a6e3a1")
        except Exception as e:
            QMessageBox.critical(self, "Export Error", str(e))

    # ── Excel import ──────────────────────────────────────────────────────────

    def _import_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Import Sequence", "", "Excel Files (*.xlsx *.xls)")
        if not path: return
        try:
            wb_in = openpyxl.load_workbook(path, data_only=True)
            rows = list(wb_in.active.iter_rows(min_row=2, values_only=True))
            self._tbl.setRowCount(0)
            for row in rows:
                if all(v is None for v in row): continue
                r = self._tbl.rowCount(); self._tbl.insertRow(r)
                for c in range(min(_N, len(row))):
                    v = row[c]
                    self._tbl.setItem(r, c, _make_item(
                        c, bool(v) if c in _CHECK else (v if v is not None else _DEFS.get(c,''))))
            # ── restore Config sheet
            cfg_map = {}
            if "Config" in wb_in.sheetnames:
                for kv in wb_in["Config"].iter_rows(min_row=2, values_only=True):
                    if kv[0] is not None:
                        cfg_map[str(kv[0])] = kv[1]
            def _g(k, default=None):
                return cfg_map.get(k, default)
            if "n_readings"   in cfg_map: self.n_readings.setValue(int(_g("n_readings")))
            if "interval_ms"  in cfg_map: self.interval_ms.setValue(int(_g("interval_ms")))
            if "settling_ms"  in cfg_map: self.settling_ms.setValue(int(_g("settling_ms")))
            if "avg_readings" in cfg_map: self.avg_readings.setChecked(bool(int(_g("avg_readings"))))
            if "gdm_en"   in cfg_map: self.gdm_en.setChecked(bool(int(_g("gdm_en"))))
            if "gdm_func" in cfg_map:
                idx = self.gdm_func.findText(str(_g("gdm_func")))
                if idx >= 0: self.gdm_func.setCurrentIndex(idx)
            if "dmm_en"   in cfg_map: self.dmm_en.setChecked(bool(int(_g("dmm_en"))))
            if "dmm_func" in cfg_map:
                idx = self.dmm_func.findText(str(_g("dmm_func")))
                if idx >= 0: self.dmm_func.setCurrentIndex(idx)
            if "itech_meas_v" in cfg_map: self.itech_meas_v.setChecked(bool(int(_g("itech_meas_v"))))
            if "itech_meas_i" in cfg_map: self.itech_meas_i.setChecked(bool(int(_g("itech_meas_i"))))
            for ch in range(1, 5):
                k = f"gpp_meas_v_{ch}"
                if k in cfg_map: self._gpp_meas_v[ch].setChecked(bool(int(_g(k))))
                k = f"gpp_meas_i_{ch}"
                if k in cfg_map: self._gpp_meas_i[ch].setChecked(bool(int(_g(k))))
            if "seq_daq_en"       in cfg_map: self.seq_daq_en.setChecked(bool(int(_g("seq_daq_en"))))
            if "seq_espec_meas_t" in cfg_map: self.seq_espec_meas_t.setChecked(bool(int(_g("seq_espec_meas_t"))))
            if "seq_espec_meas_h" in cfg_map: self.seq_espec_meas_h.setChecked(bool(int(_g("seq_espec_meas_h"))))
            if "wt_en"   in cfg_map: self.wt_en.setChecked(bool(int(_g("wt_en"))))
            if "wt_elem" in cfg_map: self.wt_elem.setValue(int(_g("wt_elem")))
            for func in self._wt_checks:
                k = f"wt_{func}"
                if k in cfg_map: self._wt_checks[func].setChecked(bool(int(_g(k))))
            if "mdo_en" in cfg_map: self.mdo_en.setChecked(bool(int(_g("mdo_en"))))
            if "mdo_ch" in cfg_map:
                idx = self.mdo_ch.findText(str(_g("mdo_ch")))
                if idx >= 0: self.mdo_ch.setCurrentIndex(idx)
            if "mdo_meas_idx" in cfg_map: self.mdo_meas.setCurrentIndex(int(_g("mdo_meas_idx")))
            # Legacy def_* keys from older exports are silently ignored
            if "on_read_fail_idx" in cfg_map:
                self.on_read_fail.setCurrentIndex(int(_g("on_read_fail_idx")))
            if "max_retries" in cfg_map:
                self.max_retries.setValue(int(_g("max_retries")))
            self._status(
                f"Imported {self._tbl.rowCount()} rows"
                + (" + config" if cfg_map else "")
                + f" from {Path(path).name}", "#f9e2af")
        except Exception as e:
            QMessageBox.critical(self, "Import Error", str(e))

    # ── Load from Sweep tab ───────────────────────────────────────────────────

    def load_steps(self, steps: list):
        if not steps: return
        if self._tbl.rowCount() > 0:
            mb = QMessageBox(self)
            mb.setWindowTitle("Load from Sweep")
            mb.setText(f"Received {len(steps)} steps from Sweep tab.")
            mb.setInformativeText("Replace existing rows or append to them?")
            btn_rep = mb.addButton("Replace", QMessageBox.ButtonRole.DestructiveRole)
            btn_app = mb.addButton("Append",  QMessageBox.ButtonRole.AcceptRole)
            mb.addButton(QMessageBox.StandardButton.Cancel)
            mb.exec()
            clicked = mb.clickedButton()
            if clicked == btn_rep:   self._tbl.setRowCount(0)
            elif clicked != btn_app: return
        for d in steps:
            self._add_row(d)
        n = len(steps)
        self._status(
            f"{'Loaded' if self._tbl.rowCount() == n else 'Appended'}"
            f" {n} steps from Sweep — edit freely then Run", "#f9e2af")
